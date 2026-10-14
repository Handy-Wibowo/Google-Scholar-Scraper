"""Local journal ranking catalogue: transactional CSV/XLSX imports and conservative matching.

No ranking data is bundled with this application. Every record here comes from a file
the user imported, and every match is either an exact ISSN match or an exact normalized
journal-title match that is not ambiguous. A missing record means "unknown", never
"unranked" — an absent match is not proof that a journal is unindexed or unaccredited.
"""

from collections.abc import Iterable, Sequence
from dataclasses import dataclass, replace
from pathlib import Path
import csv
import re
import unicodedata

from openpyxl import load_workbook

from google_scholar_scraper.models import Article, normalize_issn, normalize_issn_list


SJR_QUARTILES = ("Q1", "Q2", "Q3", "Q4")
SINTA_RANKS = ("S1", "S2", "S3", "S4", "S5", "S6")

# Accepted header aliases, already reduced through `_header_key`. Aliases cover the
# SCImago semicolon export ("Title; Issn; SJR Best Quartile; Year; Categories") and
# the SINTA journal list / template ("journal,issn,sinta_rank,year,source_url").
_COLUMN_ALIASES: dict[str, tuple[str, ...]] = {
    "title": ("title", "journal", "journaltitle", "journalname", "name", "source", "sourcetitle"),
    "issn": ("issn", "issns", "issnprint", "issnonline", "pissn", "eissn", "issnnumbers"),
    "rank": (
        "sjrbestquartile",
        "bestquartile",
        "quartile",
        "sjrquartile",
        "sinta rank",
        "sintarank",
        "rank",
        "accreditation",
        "accreditationrank",
    ),
    "year": ("year", "rankingyear", "sjryear", "sintayear", "accreditationyear", "date"),
    "category": ("categories", "category", "subjectarea", "subjectareas", "subjectcategory", "field"),
    "source_url": ("sourceurl", "url", "link", "source", "sourcelink"),
}

_SJR_RANK_RE = re.compile(r"^Q[1-4]$")
_SINTA_RANK_RE = re.compile(r"^S[1-6]$")
_NON_ALNUM_RE = re.compile(r"[^0-9a-z]+")


@dataclass(frozen=True)
class JournalRankRecord:
    title: str
    title_key: str
    issns: tuple[str, ...]
    kind: str
    rank: str
    year: int | None
    category: str = ""
    source_url: str = ""


@dataclass(frozen=True)
class JournalRank:
    """The winning record for one journal, with its provenance intact."""

    kind: str
    rank: str
    year: int | None
    category: str
    source_url: str
    matched_by: str


class JournalCatalog:
    """A local store of imported journal ranking records."""

    def __init__(self) -> None:
        self._records: list[JournalRankRecord] = []
        self._by_issn: dict[str, list[JournalRankRecord]] = {}
        self._by_title: dict[str, list[JournalRankRecord]] = {}

    def __len__(self) -> int:
        return len(self._records)

    @property
    def records(self) -> tuple[JournalRankRecord, ...]:
        return tuple(self._records)

    def import_file(self, path: str | Path, default_year: int | None = None) -> int:
        """Import ranking records from a CSV or XLSX file and add them to the catalogue.

        Imports accumulate: importing an SJR file and then a SINTA file leaves both sets of
        records in place, and importing the same source for a newer year adds that year
        alongside the older one. The import is transactional — rows are validated and
        parsed first, and the tables are only rebuilt once the whole file parsed — so a
        malformed file raises and leaves the existing records untouched.

        `default_year` is an explicit caller-supplied ranking year for exports that omit
        one (the SCImago/SINTA "latest" files). It is never inferred, and without it a
        missing year stays unknown.
        """
        file_path = Path(path)
        suffix = file_path.suffix.casefold()
        if suffix in {".xlsx", ".xlsm"}:
            headers, rows = _read_xlsx(file_path)
        elif suffix in {".csv", ".tsv", ".txt"}:
            headers, rows = _read_csv(file_path, delimiter="\t" if suffix == ".tsv" else None)
        else:
            raise ValueError(f"Unsupported ranking file type: {file_path.suffix or '(none)'}")

        incoming = _build_records(headers, rows, source_path=file_path, default_year=default_year)
        records = [*self._records, *incoming]

        by_issn: dict[str, list[JournalRankRecord]] = {}
        by_title: dict[str, list[JournalRankRecord]] = {}
        for record in records:
            for issn in record.issns:
                by_issn.setdefault(issn, []).append(record)
            if record.title_key:
                by_title.setdefault(record.title_key, []).append(record)

        self._records = records
        self._by_issn = by_issn
        self._by_title = by_title
        return len(incoming)

    def clear(self) -> None:
        self._records = []
        self._by_issn = {}
        self._by_title = {}

    def match(self, *, issns: Iterable[str] = (), journal: str = "") -> JournalRank | None:
        """Return one best rank for a journal, or None when nothing matches.

        A journal can legitimately carry an SJR quartile and a SINTA rank at the same
        time; `match_all` returns both. This convenience accessor returns the one with the
        latest year, preferring SJR on a tie, for callers that want a single label.
        """
        ranks = self.match_all(issns=issns, journal=journal)
        if not ranks:
            return None

        best = ranks[0]
        for rank in ranks[1:]:
            if _year_or_floor(rank.year) > _year_or_floor(best.year):
                best = rank
            elif _year_or_floor(rank.year) == _year_or_floor(best.year) and rank.kind == "sjr":
                best = rank
        return best

    def match_all(self, *, issns: Iterable[str] = (), journal: str = "") -> tuple[JournalRank, ...]:
        """Return at most one best record per ranking system for a journal identity.

        Every exact ISSN the caller supplies is considered together, because a journal's
        print and electronic ISSNs can carry different ranking systems — an SJR quartile on
        the e-ISSN and a SINTA rank on the p-ISSN are one journal, not a competition. Each
        system is resolved independently, so finding one never hides the other. An exact
        normalized title is only a fallback, used when no supplied ISSN matched anything.

        Within one system the latest genuinely dated record wins; if records sharing that
        latest year disagree, that system is ambiguous and is dropped rather than resolved
        by an arbitrary pick.
        """
        matched_by = ""
        candidates: list[JournalRankRecord] = []
        for issn in normalize_issn_list(issns):
            records = self._by_issn.get(issn)
            if records:
                candidates.extend(records)
                matched_by = matched_by or f"issn:{issn}"

        if not candidates:
            title_key = journal_title_key(journal)
            if not title_key:
                return ()
            candidates = list(self._by_title.get(title_key, ()))
            matched_by = "title"

        resolved = _resolve_candidates(candidates)
        return tuple(_to_rank(record, matched_by=matched_by) for record in resolved)


def journal_title_key(value: str) -> str:
    """Conservative comparison key for journal titles.

    Case, punctuation, diacritics and whitespace are removed; nothing else is. A standalone
    leading article "The" is dropped first, on the raw words, because it is the one
    journal-title variant that is unambiguously the same journal — it must be removed as a
    word, not as a prefix, or "Thermodynamics" would become "rmodynamics" and collide with
    unrelated titles. No abbreviations, no stemming, no fuzzy matching.
    """
    normalized = unicodedata.normalize("NFKD", value or "")
    decomposed = "".join(char for char in normalized if not unicodedata.combining(char))
    words = decomposed.split()
    if len(words) > 1 and words[0].casefold() == "the":
        words = words[1:]
    return _NON_ALNUM_RE.sub("", "".join(words).casefold())


def _to_rank(record: JournalRankRecord, matched_by: str) -> JournalRank:
    return JournalRank(
        kind=record.kind,
        rank=record.rank,
        year=record.year,
        category=record.category,
        source_url=record.source_url,
        matched_by=matched_by,
    )


def _resolve_candidates(candidates: Sequence[JournalRankRecord]) -> tuple[JournalRankRecord, ...]:
    """Reduce records for one journal identity to at most one winner per ranking system.

    Within a system, only the latest genuinely dated record can win: an undated record is
    never preferred over a dated one, because an undated record cannot be shown to be
    newer. If two records share that latest year but disagree on the rank, the evidence is
    contradictory and the whole system is dropped instead of picking one arbitrarily.
    """
    winners: dict[str, JournalRankRecord] = {}
    ambiguous: set[str] = set()

    for candidate in candidates:
        current = winners.get(candidate.kind)
        if current is None:
            winners[candidate.kind] = candidate
            continue

        current_year = _year_or_floor(current.year)
        candidate_year = _year_or_floor(candidate.year)

        if candidate_year > current_year:
            winners[candidate.kind] = candidate
            ambiguous.discard(candidate.kind)
        elif candidate_year == current_year and candidate.rank != current.rank:
            ambiguous.add(candidate.kind)

    return tuple(
        record for kind, record in winners.items() if kind not in ambiguous
    )


def _year_or_floor(year: int | None) -> int:
    """Sort undated records below every dated one without inventing a year for them."""
    return -1 if year is None else year


def _header_key(value: object) -> str:
    return re.sub(r"[^a-z0-9]", "", str(value or "").casefold())


def _resolve_columns(headers: Sequence[object]) -> dict[str, int]:
    keyed = {_header_key(header): index for index, header in enumerate(headers)}
    resolved: dict[str, int] = {}
    for field, aliases in _COLUMN_ALIASES.items():
        for alias in aliases:
            alias_key = _header_key(alias)
            if alias_key in keyed:
                resolved[field] = keyed[alias_key]
                break
    return resolved


def _read_csv(path: Path, delimiter: str | None) -> tuple[list[str], list[list[str]]]:
    with path.open(newline="", encoding="utf-8-sig") as handle:
        sample = handle.read(8192)
        handle.seek(0)
        if delimiter is None:
            delimiter = ";" if sample.count(";") > sample.count(",") else ","
        reader = csv.reader(handle, delimiter=delimiter)
        rows = [row for row in reader]

    if not rows:
        raise ValueError(f"Ranking file {path.name} is empty.")
    return rows[0], rows[1:]


def _read_xlsx(path: Path) -> tuple[list[str], list[list[str]]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        worksheet = workbook.active
        rows = [list(row) for row in worksheet.iter_rows(values_only=True)]
    finally:
        workbook.close()

    if not rows:
        raise ValueError(f"Ranking file {path.name} is empty.")
    headers = ["" if value is None else str(value) for value in rows[0]]
    return headers, [[("" if value is None else value) for value in row] for row in rows[1:]]


def _build_records(
    headers: Sequence[object],
    rows: Sequence[Sequence[object]],
    *,
    source_path: Path,
    default_year: int | None = None,
) -> list[JournalRankRecord]:
    columns = _resolve_columns(headers)
    if "title" not in columns and "issn" not in columns:
        raise ValueError(
            f"Ranking file {source_path.name} has no recognizable journal title or ISSN column. "
            f"Headers found: {', '.join(str(header) for header in headers if str(header).strip())}"
        )
    if "rank" not in columns:
        raise ValueError(f"Ranking file {source_path.name} has no recognizable rank/quartile column.")

    records: list[JournalRankRecord] = []
    for row_number, row in enumerate(rows, start=2):
        if not any(str(value).strip() for value in row):
            continue

        title = str(_cell(row, columns.get("title"))).strip()
        issns = normalize_issn_list(_split_issns(_cell(row, columns.get("issn"))))
        if not title and not issns:
            continue

        raw_rank = _cell(row, columns.get("rank"))
        kind, rank = _classify_rank(raw_rank)
        if not rank:
            continue

        year = _parse_year(_cell(row, columns.get("year")))
        if year is None:
            year = default_year

        source_url = str(_cell(row, columns.get("source_url"))).strip()
        records.append(
            JournalRankRecord(
                title=title,
                title_key=journal_title_key(title),
                issns=issns,
                kind=kind,
                rank=rank,
                year=year,
                category=str(_cell(row, columns.get("category"))).strip(),
                source_url=source_url,
            )
        )

    if not records:
        raise ValueError(
            f"Ranking file {source_path.name} contained no valid ranking rows "
            f"(ranks must be Q1-Q4 or S1-S6)."
        )
    return records


def _cell(row: Sequence[object], index: int | None) -> object:
    if index is None or index >= len(row):
        return ""
    return row[index]


def _split_issns(value: object) -> list[str]:
    return [part for part in re.split(r"[;,/|\s]+", str(value or "")) if part]


def _classify_rank(value: object) -> tuple[str, str]:
    token = re.sub(r"\s+", "", str(value or "")).upper()
    if _SJR_RANK_RE.match(token):
        return "sjr", token
    if _SINTA_RANK_RE.match(token):
        return "sinta", token
    return "", ""


def _parse_year(value: object) -> int | None:
    match = re.search(r"(1[5-9]\d\d|2[01]\d\d)", str(value or ""))
    return int(match.group(1)) if match else None


def apply_rank(article: Article, rank: JournalRank | None) -> Article:
    """Attach a matched rank to an article, filling only that ranking system's fields."""
    if rank is None:
        return article

    if rank.kind == "sjr":
        return replace(
            article,
            sjr_quartile=rank.rank,
            sjr_year=rank.year,
            sjr_category=rank.category,
            sjr_source_url=rank.source_url,
        )
    return replace(
        article,
        sinta_rank=rank.rank,
        sinta_year=rank.year,
        sinta_source_url=rank.source_url,
    )


def catalog_matches(article: Article, catalog: JournalCatalog | None) -> tuple[JournalRank | None, JournalRank | None]:
    """Return the (SJR, SINTA) records for an article's known journal identity.

    The most specific identity that yields anything wins, so an ISSN-verified match is
    never downgraded by a title-only match from a vaguer identity.
    """
    if catalog is None:
        return None, None

    for issns, journal in _journal_identities(article):
        ranks = catalog.match_all(issns=issns, journal=journal)
        if not ranks:
            continue

        sjr = next((rank for rank in ranks if rank.kind == "sjr"), None)
        sinta = next((rank for rank in ranks if rank.kind == "sinta"), None)
        if sjr is not None or sinta is not None:
            return sjr, sinta

    return None, None


def _journal_identities(article: Article) -> list[tuple[tuple[str, ...], str]]:
    """Candidate journal identities, most specific first.

    Only journal names and ISSNs that were themselves verified count here: the parsed
    journal field and a normalized DOI. The Scholar citation line is never used — it
    mixes authors, venue and year, and guessing a journal from it would fabricate an
    identity the source never asserted.
    """
    identities: list[tuple[tuple[str, ...], str]] = [(article.issns, article.journal)]

    from_doi = _issn_from_doi(article.doi)
    if from_doi:
        identities.insert(0, ((from_doi,), ""))

    return identities


def _issn_from_doi(doi: str) -> str:
    match = re.search(r"10\.\d{4,9}/(\d{4}-\d{3}[\dX])", (doi or "").upper())
    return normalize_issn(match.group(1)) if match else ""
