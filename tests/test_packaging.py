import ast
import csv
import os
import runpy
import sys
import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

import google_scholar_scraper
from google_scholar_scraper.app import main, run_packaging_smoke


PROJECT_ROOT = Path(__file__).resolve().parents[1]


class TclTkPackagingTests(unittest.TestCase):
    def test_spec_removes_stock_hook_from_actual_script_table(self) -> None:
        spec = PROJECT_ROOT / "packaging" / "pyinstaller" / "GoogleScholarScraper.spec"
        assignments = [node for node in ast.parse(spec.read_text(encoding="utf-8")).body
                       if isinstance(node, ast.Assign)
                       and any(isinstance(target, ast.Attribute)
                               and isinstance(target.value, ast.Name)
                               and target.value.id == "a" and target.attr == "scripts"
                               for target in node.targets)]
        self.assertEqual(len(assignments), 1)
        custom = ("pyi_rth_tkinter_embedded", "custom.py", "PYSOURCE")
        entrypoint = ("__main__", "__main__.py", "PYSOURCE")
        analysis = SimpleNamespace(scripts=[custom, ("pyi_rth__tkinter", "stock.py", "PYSOURCE"), entrypoint])
        module = ast.fix_missing_locations(ast.Module(body=assignments, type_ignores=[]))
        exec(compile(module, str(spec), "exec"), {"a": analysis})
        self.assertEqual(analysis.scripts, [custom, entrypoint])

    def test_hook_points_at_real_bundled_script_directories(self) -> None:
        hook = PROJECT_ROOT / "packaging" / "pyinstaller" / "pyi_rth_tkinter_embedded.py"
        with tempfile.TemporaryDirectory() as folder:
            for name in ("_tcl_data", "_tk_data"):
                (Path(folder) / name).mkdir()
            with patch.object(sys, "_MEIPASS", folder, create=True), patch.dict(os.environ, {}, clear=True):
                runpy.run_path(str(hook))
                self.assertEqual(os.environ["TCL_LIBRARY"], str(Path(folder) / "_tcl_data"))
                self.assertEqual(os.environ["TK_LIBRARY"], str(Path(folder) / "_tk_data"))

    def test_embedded_hook_removes_external_library_overrides(self) -> None:
        hook = PROJECT_ROOT / "packaging" / "pyinstaller" / "pyi_rth_tkinter_embedded.py"
        with tempfile.TemporaryDirectory() as folder:
            with patch.object(sys, "_MEIPASS", folder, create=True), patch.dict(
                os.environ, {"TCL_LIBRARY": "external-tcl", "TK_LIBRARY": "external-tk"}, clear=True
            ):
                runpy.run_path(str(hook))
                self.assertNotIn("TCL_LIBRARY", os.environ)
                self.assertNotIn("TK_LIBRARY", os.environ)
                self.assertFalse((Path(folder) / "_tcl_data").exists())
                self.assertFalse((Path(folder) / "_tk_data").exists())


class PackagingMetadataTests(unittest.TestCase):
    def test_package_version_matches_project_metadata(self) -> None:
        pyproject = (PROJECT_ROOT / "pyproject.toml").read_text(encoding="utf-8")

        self.assertEqual(google_scholar_scraper.__version__, "2.2.0")
        self.assertIn('version = "2.2.0"', pyproject)
        self.assertIn('requires-python = ">=3.10"', pyproject)
        self.assertIn('"requests>=2.33,<3"', pyproject)

    def test_gui_entrypoint_is_import_safe(self) -> None:
        self.assertTrue(callable(main))

    def test_packaging_smoke_writes_excel_and_csv_without_starting_gui(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            run_packaging_smoke(Path(temp_dir))

            self.assertTrue((Path(temp_dir) / "packaging-smoke.xlsx").exists())
            csv_path = Path(temp_dir) / "packaging-smoke.csv"
            self.assertTrue(csv_path.exists())
            with csv_path.open(newline="", encoding="utf-8-sig") as handle:
                rows = list(csv.DictReader(handle))
            self.assertEqual(rows[0]["Title"], "Unicode Smoke: cafe, alpha + beta, 医学 AI")
            self.assertEqual(rows[0]["Relevance Score"], "88.8")
            self.assertEqual(rows[1]["Relevance Score"], "")

    def test_packaging_smoke_exercises_every_export_column_offline(self) -> None:
        from openpyxl import load_workbook

        from google_scholar_scraper.exporters import EXPORT_COLUMNS

        with tempfile.TemporaryDirectory() as temp_dir:
            run_packaging_smoke(Path(temp_dir))

            csv_path = Path(temp_dir) / "packaging-smoke.csv"
            with csv_path.open(newline="", encoding="utf-8-sig") as handle:
                rows = list(csv.DictReader(handle))

            # The smoke fixture must cover the whole export surface, so a column added
            # to the model cannot slip out of the packaged smoke check unnoticed.
            self.assertEqual(list(rows[0].keys()), EXPORT_COLUMNS)

            ranked = rows[0]
            self.assertEqual(ranked["DOI"], "10.1000/smoke.2024.001")
            self.assertEqual(ranked["Journal"], "Journal of Packaging Smoke")
            self.assertEqual(ranked["ISSNs"], "1234-5679; 1234-5687")
            self.assertEqual(ranked["Publication Year"], "2024")
            self.assertEqual(ranked["SJR Quartile"], "Q1")
            self.assertEqual(ranked["SJR Year"], "2024")
            self.assertEqual(ranked["SJR Category"], "Library and Information Sciences")
            self.assertTrue(ranked["SJR Source"].startswith("https://www.scimagojr.com/"))
            self.assertEqual(ranked["SINTA Rank"], "S2")
            self.assertEqual(ranked["SINTA Year"], "2024")
            self.assertTrue(ranked["SINTA Source"].startswith("https://sinta.kemdikbud.go.id/"))
            # The fixture never checks or downloads a PDF, so it must report the
            # neutral state instead of claiming availability.
            self.assertEqual(ranked["PDF Status"], "unknown")
            self.assertEqual(ranked["PDF URL"], "https://example.edu/smoke.pdf")
            self.assertEqual(ranked["PDF Source"], "https://example.edu/smoke")
            self.assertIn("no public-PDF check was performed", ranked["PDF Detail"])
            self.assertEqual(ranked["PDF Local Path"], "")

            unranked = rows[1]
            for column in ("DOI", "Journal", "ISSNs", "SJR Quartile", "SINTA Rank", "PDF URL"):
                self.assertEqual(unranked[column], "", column)
            # A record that was never looked up reports the neutral default, not a
            # claim that no public PDF exists.
            self.assertEqual(unranked["PDF Status"], "unknown")

            workbook = load_workbook(Path(temp_dir) / "packaging-smoke.xlsx")
            worksheet = workbook.active
            header = [cell.value for cell in worksheet[1]]
            self.assertEqual(header, EXPORT_COLUMNS)
            excel_row = {column: worksheet.cell(row=2, column=index + 1).value for index, column in enumerate(header)}
            self.assertEqual(excel_row["SJR Quartile"], "Q1")
            self.assertEqual(excel_row["SINTA Rank"], "S2")
            self.assertEqual(excel_row["PDF Status"], "unknown")
            self.assertEqual(excel_row["Publication Year"], 2024)

    def test_pyinstaller_spec_uses_windowed_onedir_identity_without_duplicate_notice_datas(self) -> None:
        spec = (PROJECT_ROOT / "packaging" / "pyinstaller" / "GoogleScholarScraper.spec").read_text(encoding="utf-8")

        self.assertIn('entrypoint = src_path / "google_scholar_scraper" / "__main__.py"', spec)
        self.assertIn("console=False", spec)
        self.assertIn('name="Google-Scholar-Scraper-v2.2.0"', spec)
        self.assertIn("datas=[]", spec)
        self.assertNotIn("distribution_documents", spec)
        self.assertIn('excludes=["tests", *excluded_optional_modules]', spec)
        self.assertIn('"h2"', spec)
        self.assertIn('"pandas"', spec)
        self.assertIn('"torch"', spec)

    def test_build_script_uses_expected_artifact_names_and_license_collection(self) -> None:
        script = (PROJECT_ROOT / "scripts" / "build_windows.ps1").read_text(encoding="utf-8")

        self.assertIn("Google-Scholar-Scraper-v$Version-Portable-Windows-x64.zip", script)
        self.assertIn("Google-Scholar-Scraper-v$Version-Setup-Windows-x64.exe", script)
        self.assertIn("Refusing to remove path outside repository", script)
        self.assertIn('"build", "dist"', script)
        self.assertIn("System.Security.Cryptography.SHA256", script)
        self.assertIn("Collect-ThirdPartyLicenses", script)
        self.assertIn("Verify-DistributionCompliance", script)
        self.assertIn("collect_third_party_licenses.py", script)
        self.assertIn("third_party_licenses", script)
        self.assertNotIn("installer/$InstallerName", script)

    def test_third_party_license_collector_covers_runtime_and_packaging_components(self) -> None:
        collector = (PROJECT_ROOT / "scripts" / "collect_third_party_licenses.py").read_text(encoding="utf-8")

        for distribution_name in (
            "beautifulsoup4",
            "soupsieve",
            "openpyxl",
            "et-xmlfile",
            "requests",
            "urllib3",
            "charset-normalizer",
            "idna",
            "certifi",
            "typing-extensions",
            "pyinstaller",
            "pyinstaller-hooks-contrib",
        ):
            self.assertIn(f'"{distribution_name}"', collector)
        self.assertIn("copy_python_license", collector)
        self.assertIn("copy_tcl_tk_licenses", collector)
        self.assertIn("MANIFEST.txt", collector)

    def test_inno_setup_definition_uses_per_user_installer_metadata(self) -> None:
        iss = (PROJECT_ROOT / "installer" / "GoogleScholarScraper.iss").read_text(encoding="utf-8")

        self.assertIn('#define MyAppVersion "2.2.0"', iss)
        self.assertIn('#define MyAppPublisher "Mahdi Navaei"', iss)
        self.assertIn("PrivilegesRequired=lowest", iss)
        self.assertIn("Google-Scholar-Scraper-v2.2.0-Setup-Windows-x64", iss)

    def test_release_licensing_notice_and_release_note_files_are_present(self) -> None:
        license_text = (PROJECT_ROOT / "LICENSE").read_text(encoding="utf-8")
        notice = (PROJECT_ROOT / "NOTICE").read_text(encoding="utf-8")
        commercial = (PROJECT_ROOT / "COMMERCIAL_LICENSE.md").read_text(encoding="utf-8")
        third_party = (PROJECT_ROOT / "THIRD_PARTY_NOTICES.txt").read_text(encoding="utf-8")
        release_notes = (PROJECT_ROOT / "docs" / "RELEASE_NOTES_V2.0.1.md").read_text(encoding="utf-8")
        current_notes = (PROJECT_ROOT / "docs" / "RELEASE_NOTES_V2.2.0.md").read_text(encoding="utf-8")
        changelog = (PROJECT_ROOT / "CHANGELOG.md").read_text(encoding="utf-8")

        self.assertIn("PolyForm Noncommercial License 1.0.0", license_text)
        self.assertIn("Mahdi Navaei", notice)
        self.assertIn("Commercial use requires a separate written commercial license", commercial)
        self.assertIn("beautifulsoup4 4.14.3", third_party)
        self.assertIn("requests 2.33.0", third_party)
        self.assertIn("urllib3 2.7.0", third_party)
        self.assertIn("idna 3.15", third_party)
        self.assertIn("Google Scholar Scraper V2.0.1", release_notes)
        self.assertIn("Google Scholar Scraper V2.2.0", current_notes)
        self.assertIn("## 2.2.0", changelog)

    def test_release_notes_state_the_ranking_and_pdf_caveats(self) -> None:
        notes = (PROJECT_ROOT / "docs" / "RELEASE_NOTES_V2.2.0.md").read_text(encoding="utf-8")

        # The release must not over-promise: no bundled dataset, no automatic
        # ranking feed, no access-control bypass, and no claim that a candidate
        # PDF link proves a reuse licence.
        self.assertIn("No bundled ranking dataset", notes)
        self.assertIn("does not supply SJR or SINTA ranking labels", notes)
        self.assertIn("not an official Scopus CiteScore quartile", notes)
        self.assertIn("unavailable", notes)
        self.assertIn("unknown", notes)
        self.assertIn("Candidate PDF metadata does not prove a reuse licence", notes)

    def test_distributions_ship_user_documentation_and_import_templates(self) -> None:
        script = (PROJECT_ROOT / "scripts" / "build_windows.ps1").read_text(encoding="utf-8")
        workflow = (PROJECT_ROOT / ".github" / "workflows" / "build-windows.yml").read_text(encoding="utf-8")
        readme = (PROJECT_ROOT / "README.md").read_text(encoding="utf-8")

        for build_source in (script, workflow):
            self.assertIn("ranking-imports.md", build_source)
            self.assertIn("ranking_templates", build_source)
            self.assertIn("README.md", build_source)
            # Fail closed: a missing guide or template set must break the build
            # rather than silently produce a distribution without them.
            self.assertIn("No header-only ranking CSV templates found in", build_source)
            self.assertIn("Required user documentation missing from distribution:", build_source)
            self.assertIn("Distribution ships no ranking import templates.", build_source)

        self.assertIn("docs/ranking-imports.md", readme)
        self.assertIn("does **not** bundle a ranking dataset", readme)

    def test_docs_link_to_the_intended_v2_2_0_release_and_keep_caveats(self) -> None:
        readme = (PROJECT_ROOT / "README.md").read_text(encoding="utf-8")
        notes = (PROJECT_ROOT / "docs" / "RELEASE_NOTES_V2.2.0.md").read_text(encoding="utf-8")
        changelog = (PROJECT_ROOT / "CHANGELOG.md").read_text(encoding="utf-8")

        # V2.2.0 is the current normal release of the fork: the docs point at the
        # intended tag and asset URLs that go live once CI and Windows installer
        # checks pass on the release commit and the release is published. Asserting
        # these keeps coverage of the intentional release URLs instead of disabling it.
        tag_url = "https://github.com/Handy-Wibowo/Google-Scholar-Scraper/releases/tag/v2.2.0"
        installer_url = "https://github.com/Handy-Wibowo/Google-Scholar-Scraper/releases/download/v2.2.0/Google-Scholar-Scraper-v2.2.0-Setup-Windows-x64.exe"
        portable_url = "https://github.com/Handy-Wibowo/Google-Scholar-Scraper/releases/download/v2.2.0/Google-Scholar-Scraper-v2.2.0-Portable-Windows-x64.zip"
        checksums_url = "https://github.com/Handy-Wibowo/Google-Scholar-Scraper/releases/download/v2.2.0/SHA256SUMS.txt"

        for document in (readme, notes):
            self.assertIn(tag_url, document)
        self.assertIn(installer_url, readme)
        self.assertIn(portable_url, readme)
        self.assertIn(checksums_url, readme)
        self.assertIn(tag_url, changelog)

        # The V2.2.0 docs no longer label it an unpublished feature-branch preview.
        # The historical V2.1.0 changelog entry keeps its own wording unchanged.
        self.assertNotIn("unpublished feature-branch preview", readme)
        self.assertNotIn("unpublished feature-branch preview", notes)
        self.assertNotIn("Unpublished feature-branch preview on", changelog)

        # Upstream V2.0.1 stays the historical published upstream release, with
        # attribution preserved.
        self.assertIn("releases/tag/v2.0.1", readme)
        self.assertIn("MahdiNavaei/Google-Scholar-Scraper", readme)
        self.assertIn("Handy-Wibowo/Google-Scholar-Scraper", readme)
        self.assertIn("Mahdi Navaei", readme)
        self.assertIn("not been proposed to, reviewed by, or endorsed by the upstream", readme)

        # Release caveats survive the release framing.
        self.assertIn("does **not** bundle a ranking dataset", readme)
        self.assertIn("official Scopus CiteScore quartile", readme)
        self.assertIn("not an official Scopus CiteScore quartile", notes)

    def test_windows_workflow_builds_installer_collects_licenses_and_uploads_artifacts_without_release_publish(self) -> None:
        workflow = (PROJECT_ROOT / ".github" / "workflows" / "build-windows.yml").read_text(encoding="utf-8")

        self.assertIn("workflow_dispatch", workflow)
        self.assertIn("pull_request", workflow)
        self.assertIn("choco install innosetup", workflow)
        self.assertIn("Build Windows Installer", workflow)
        self.assertIn("Collect exact third-party licenses", workflow)
        self.assertIn("Verify packaged executable and distribution compliance", workflow)
        self.assertIn("Installer install-launch-export-uninstall smoke", workflow)
        self.assertIn("unins000.exe", workflow)
        self.assertIn("constraints/release-2.2.0.txt", workflow)
        self.assertIn("third_party_licenses", workflow)
        self.assertIn("actions/upload-artifact@v4", workflow)
        self.assertNotIn("softprops/action-gh-release", workflow)

    def test_packaged_smoke_verifies_the_new_export_columns(self) -> None:
        workflow = (PROJECT_ROOT / ".github" / "workflows" / "build-windows.yml").read_text(encoding="utf-8")

        # The packaged smoke must assert the columns, not merely that two files
        # exist — otherwise a packaging regression could ship a truncated export.
        self.assertIn("Packaged smoke CSV is missing the", workflow)
        for column in ("SJR Quartile", "SINTA Rank", "PDF Status", "PDF Local Path"):
            self.assertIn(f'"{column}"', workflow)
        self.assertIn("Installed user documentation missing:", workflow)
        self.assertIn("Installed distribution ships no ranking import templates.", workflow)

    def test_release_constraints_pin_expected_dependency_versions(self) -> None:
        constraints = (PROJECT_ROOT / "constraints" / "release-2.2.0.txt").read_text(encoding="utf-8")

        self.assertIn("beautifulsoup4==4.14.3", constraints)
        self.assertIn("requests==2.33.0", constraints)
        self.assertIn("urllib3==2.7.0", constraints)
        self.assertIn("idna==3.15", constraints)
        self.assertIn("pyinstaller==6.20.0", constraints)

    def test_release_constraints_reuse_the_previous_release_pins_verbatim(self) -> None:
        def pins(name: str) -> list[str]:
            text = (PROJECT_ROOT / "constraints" / name).read_text(encoding="utf-8")
            return [line.strip() for line in text.splitlines() if line.strip() and not line.startswith("#")]

        self.assertEqual(pins("release-2.2.0.txt"), pins("release-2.0.1.txt"))


if __name__ == "__main__":
    unittest.main()
